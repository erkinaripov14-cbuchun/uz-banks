"""Natijalarni JSON va Excel (report.xlsx) ga yozish."""
import json
from pathlib import Path

SECTION_TITLES = {"credits": "Kreditlar", "deposits": "Omonatlar"}


def _rng(a, b, fmt="{:g}"):
    if a is None and b is None:
        return ""
    if a == b or b is None:
        return fmt.format(a)
    return f"{fmt.format(a)} – {fmt.format(b)}"


def _money(v):
    if v is None:
        return None
    for d, s in ((1e9, " mlrd"), (1e6, " mln"), (1e3, " ming")):
        if v >= d:
            return f"{v / d:g}{s}"
    return f"{v:g}"


def _amount(a, b):
    if a is None:
        return ""
    return _money(a) if a == b else f"{_money(a)} – {_money(b)}"


def _rate_context(rates, n=3):
    return " • ".join(dict.fromkeys(r["context"] for r in rates[:n]))


def write_json(results, out_dir: Path, sections):
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "all.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    for s in sections:
        flat = [{"bank": r["bank"], "bank_id": r["bank_id"], **{k: v for k, v in p.items() if k != "text"}}
                for r in results for p in r.get(s, [])]
        (out_dir / f"{s}.json").write_text(json.dumps(flat, ensure_ascii=False, indent=2), encoding="utf-8")


def write_excel(results, path: Path, sections):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E78")

    def sheet(title, header, rows, widths):
        ws = wb.create_sheet(title[:31])
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
            link = row[-1]
            if link.value and str(link.value).startswith("http"):
                link.hyperlink, link.style = link.value, "Hyperlink"
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions

    for s in sections:
        name = SECTION_TITLES.get(s, s)
        # 1) Har bir mahsulot sahifasi
        rows = []
        for r in results:
            for p in r.get(s, []):
                rows.append([
                    r["bank"], p["title"], "Ro'yxat" if p["kind"] == "listing" else "Mahsulot",
                    _rng(p["rate_min"], p["rate_max"]), _rng(p["term_min_months"], p["term_max_months"]),
                    _amount(p["amount_min"], p["amount_max"]), ", ".join(p["currencies"]),
                    _rate_context(p["rates"]), p["url"],
                ])
        sheet(name, ["Bank", "Mahsulot", "Sahifa turi", "Foiz, % yillik", "Muddat, oy",
                     "Summa (so'm)", "Valyuta", "Foiz haqida matn (saytdan)", "Havola"],
              rows, [24, 40, 12, 14, 14, 20, 12, 70, 50])

        # 2) Ro'yxat sahifalaridagi mahsulot kartochkalari
        rows = []
        for r in results:
            for p in r.get(s, []):
                for c in p.get("cards", []):
                    rows.append([
                        r["bank"], c["text"].split("\n")[0][:120],
                        _rng(c["rate_min"], c["rate_max"]), _rng(c["term_min_months"], c["term_max_months"]),
                        _amount(c["amount_min"], c["amount_max"]), ", ".join(c["currencies"]),
                        c["text"][:600], c["url"],
                    ])
        sheet(f"{name} (kartochkalar)", ["Bank", "Mahsulot", "Foiz, %", "Muddat, oy", "Summa (so'm)",
                                         "Valyuta", "Kartochka matni", "Havola"],
              rows, [24, 40, 14, 14, 20, 12, 70, 50])

    # 3) Bank bo'yicha xulosa va xatolar
    rows = []
    for r in results:
        rows.append([r["bank"], *[len(r.get(s, [])) for s in sections],
                     "\n".join(r["errors"])[:2000], r["site"]])
    sheet("Banklar", ["Bank", *[SECTION_TITLES.get(s, s) + " (sahifa)" for s in sections],
                      "Xatolar", "Havola"], rows, [28, *[14] * len(sections), 80, 36])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
